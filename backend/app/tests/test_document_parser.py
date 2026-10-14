from io import BytesIO

import fitz
from docx import Document

from app.services.document_parser import parse_docx, parse_pdf


def _make_pdf_bytes() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "First page text")
    second_page = doc.new_page()
    second_page.insert_text((72, 72), "Second page text")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _make_docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Introduction", level=1)
    document.add_paragraph("This is a paragraph.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header A"
    table.cell(0, 1).text = "Header B"
    table.cell(1, 0).text = "Row 1, Col 1"
    table.cell(1, 1).text = "Row 1, Col 2"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_parse_pdf_extracts_two_pages() -> None:
    result = parse_pdf(_make_pdf_bytes())

    assert result.file_type == "pdf"
    assert len(result.blocks) >= 2
    assert result.blocks[0].page_number == 1
    assert result.blocks[1].page_number == 2
    assert all(
        "First page text" in block.text or "Second page text" in block.text
        for block in result.blocks
    )


def test_parse_docx_keeps_headings_and_tables() -> None:
    result = parse_docx(_make_docx_bytes())

    assert result.file_type == "docx"
    assert any(block.block_type == "heading" for block in result.blocks)
    assert any(block.block_type == "table" for block in result.blocks)
    assert all(block.page_number is None for block in result.blocks)


def test_parse_docx_preserves_paragraph_and_table_body_order() -> None:
    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Table content"
    document.add_paragraph("After table")
    buffer = BytesIO()
    document.save(buffer)

    result = parse_docx(buffer.getvalue())

    assert [block.text for block in result.blocks] == [
        "Before table",
        "Table content",
        "After table",
    ]
